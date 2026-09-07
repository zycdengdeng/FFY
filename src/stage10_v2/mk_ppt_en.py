# -*- coding: utf-8 -*-
"""本周英文汇报 PPT（10 页）。所有数字取自 outputs/ 下的分析产物。"""
import os, json
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN
U = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
F = f"{U}/outputs/ppt_en"; A = f"{U}/outputs/anim_b"
W, H = 13.333, 7.5
INK = C(0x1A, 0x1A, 0x1A); BLU = C(0x1B, 0x6C, 0xA8); GRY = C(0x5A, 0x5A, 0x5A)
GRN = C(0x2E, 0x7D, 0x5B); RED = C(0xC0, 0x39, 0x2B); LGT = C(0xF2, 0xF4, 0xF7)
prs = Presentation(); prs.slide_width = I(W); prs.slide_height = I(H)
BLANK = prs.slide_layouts[6]

def tb(sl, x, y, w, h, txt, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, space=1.0):
    t = sl.shapes.add_textbox(I(x), I(y), I(w), I(h)); f = t.text_frame
    f.word_wrap = True; f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    for i, line in enumerate(txt.split("\n")):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.text = line; p.alignment = align
        p.line_spacing = space
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
            r.font.name = "Calibri"
    return t

def band(sl, y, h, col):
    from pptx.enum.shapes import MSO_SHAPE
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(0), I(y), I(W), I(h))
    s.fill.solid(); s.fill.fore_color.rgb = col; s.line.fill.background(); s.shadow.inherit = False
    return s

def pic_fit(sl, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size; ar = iw/ih
    if w/h > ar: ww, hh = h*ar, h
    else:        ww, hh = w, w/ar
    return sl.shapes.add_picture(path, I(x+(w-ww)/2), I(y+(h-hh)/2), I(ww), I(hh))

def slide(title, kicker=None):
    sl = prs.slides.add_slide(BLANK)
    tb(sl, .48, .30, W-.96, .55, title, 26, True, INK)
    if kicker: tb(sl, .48, .90, W-.96, .34, kicker, 14.5, False, GRY)
    return sl

def takeaway(sl, txt, col=BLU):
    band(sl, H-.86, .015, C(0xDD, 0xDD, 0xDD))
    tb(sl, .48, H-.72, W-.96, .52, txt, 15.5, True, col)

def notes(sl, txt):
    sl.notes_slide.notes_text_frame.text = txt

# ============ 1 title ============
sl = prs.slides.add_slide(BLANK)
band(sl, 0, .28, BLU)
tb(sl, .9, 1.55, W-1.8, 1.3,
   "Can 8,705 birds tell us how long a landing gear should be?", 40, True, INK)
tb(sl, .9, 3.05, W-1.8, .9,
   "A phylogenetically controlled leg-length law, and what happens when a machine\n"
   "is allowed to disagree with it", 20, False, GRY, space=1.25)
band(sl, 4.25, .02, C(0xCC, 0xCC, 0xCC))
tb(sl, .9, 4.55, 7.4, 1.6,
   "Zihan Wang\nSupervisors:  Yifan Zhao  ·  K. Jimmy Hsia\nSJTU – EIT joint PhD programme",
   16, False, INK, space=1.35)
tb(sl, 8.6, 4.55, 3.9, 1.6, "Group meeting\n7 September 2026", 16, False, GRY, space=1.35)
notes(sl, "15 minutes. The talk is about what we are doing, not everything we have done. "
          "Two halves: (1) a law we extracted from birds and how hard we tried to break it; "
          "(2) the machine we hand that law to, and the one number where the machine "
          "disagrees with biology.")

# ============ 2 dataset ============
sl = slide("Step 0 — the table we had to build before anything else",
           "Nothing new was measured. The contribution is the join — and the phylogeny attached to every row.")
pic_fit(sl, f"{F}/en_F7_dataset.png", .55, 1.35, W-1.1, 5.0)
takeaway(sl, "8,705 species · 188 families · 35 orders, each carrying morphology, "
             "ecology and a position on the avian tree.")
notes(sl, "AVONET gives the morphology, EltonTraits the ecology, BirdTree the ancestry. "
          "The names differ between BirdLife and BirdTree taxonomies, so a crosswalk is needed — "
          "98.6% of our species matched. Flightless taxa are removed: an ostrich would dominate "
          "any leg-length regression for the wrong reason. HWI is the hand-wing index, Kipp's "
          "distance over wing length: a pointed, high-aspect wing gives a high HWI and means "
          "cheap, efficient, long-distance flight. It is the standard proxy in this literature.")

# ============ 3 law 1 ============
sl = slide("Step 1 — fit the obvious law, then throw the law away",
           "Classical allometry. What we keep is the part it does NOT explain: the residual u.")
pic_fit(sl, f"{F}/en_F1_allometry.png", .55, 1.35, W-1.1, 4.95)
takeaway(sl, "Our design prior is the phylogenetically corrected waterbird fit:  b = 0.365,  σ = 0.081 dex.  "
             "Everything after this slide is about the residual u.")
notes(sl, "Left: 8,705 species. The orange dashed line is what you get if you regress naively; "
          "the green line is the clade we actually take the prior from — waterbirds, the closest "
          "ecological analogue to a heavy VTOL that lands on soft, uneven ground — fitted with the "
          "phylogeny in the covariance. b = 0.365 means the leg grows more slowly than the body; "
          "geometric similarity would give 1/3, so this is close to isometric with a slight excess. "
          "Right: subtract the line, divide by the scatter, and you get u, a z-score of leg length "
          "for a bird of that mass. Negative u = shorter leg than its mass predicts. "
          "u is the dependent variable for the rest of the biology.")

# ============ 4 law 2 ============
sl = slide("Step 2 — one ecological variable earns its place in the law",
           "HWI = hand-wing index = Kipp's distance / wing length. High HWI = pointed wing = cheap, efficient flight.")
pic_fit(sl, f"{F}/en_F2_hwi.png", .55, 1.35, W-1.1, 4.95)
takeaway(sl, "R² 0.671 → 0.793 from a single term.  A sparrow-to-waterbird change in wing shape "
             "predicts a 31% shorter leg at the same body mass.", GRN)
notes(sl, "This is the fit we had not done before: HWI directly in the leg-length formula, not "
          "against the residual. One extra term takes R² from 0.67 to 0.79 — it absorbs 37% of the "
          "variance the mass-only law left behind. The coefficient is −0.0061 per HWI unit, t = −71. "
          "Read it as a trade-off: a bird that invests in efficient flight pays for it by giving up "
          "leg. That is exactly the trade a landing gear designer faces — stroke length versus "
          "carried mass.")

# ============ 5 pgls ============
sl = slide("Step 3 — the objection a biologist would raise first",
           "Closely related species are not independent data points. Phylogenetic GLS, 200 trees × 2 backbones.")
pic_fit(sl, f"{F}/en_F3_pgls.png", .40, 1.40, W-.80, 4.55)
takeaway(sl, "Phylogeny eats 65% of the naive effect — and the remaining 35% survives every tree "
             "(p < 1e-56, λ = 0.93).")
notes(sl, "This is the slide a biologist would demand. Ordinary least squares assumes independent "
          "residuals; on a phylogeny they are not — two swifts resemble each other because of a "
          "shared ancestor, not because each independently solved the problem. PGLS puts the "
          "expected covariance from the tree into the error term. λ = 0.93 says the trait is almost "
          "fully phylogenetically structured, so the correction is severe: the coefficient drops "
          "from −0.078 to −0.027. But it never touches zero, on any of the 200 trees. "
          "We report a shrunken, defensible effect rather than the flattering one.")

# ============ 6 replication ============
sl = slide("Step 4 — make the tree argue with itself",
           "If this were one ancient accident, it would not have to happen again inside every family.")
pic_fit(sl, f"{F}/en_F46_replication.png", .55, 1.28, W-1.1, 5.05)
takeaway(sl, "77 of 96 families and 251 of 419 genera repeat it internally; "
             "24 non-nested clades shifted leg length on their own, and how much they shifted "
             "tracks their flight efficiency (r = −0.59).", GRN)
notes(sl, "Panel (a): inside a single family, remove mass, split into better- and worse-flying "
          "halves, compare. Same ancestry, same lifestyle — the only thing left is flight "
          "efficiency. 80% of families go the predicted way; chance would give 50%. Tighten to "
          "genus and it still holds at 60%. Panel (b): 24 clades that are mutually non-nested, so "
          "each shift is a separate evolutionary event, not one event counted 24 times. The clades "
          "that shortened most are the ones that fly best. This is the result I would put in the "
          "paper's abstract.")

# ============ 7 leg use ============
sl = slide("Step 5 — a function check, from a completely separate dataset",
           "If the law is real, leg length should also track how the leg is used — not only how the wing is shaped.")
pic_fit(sl, f"{F}/en_F5_forstrat.png", .45, 1.40, W-.90, 4.60)
takeaway(sl, "Aerial feeders carry the shortest legs (−2.6 SD), ground feeders the longest (+0.9 SD); "
             "the sign of the partial correlation flips exactly where function says it should.")
notes(sl, "EltonTraits gives, for every species, the percentage of foraging done in each stratum. "
          "Birds that feed in the air barely use the leg and have almost none; birds that walk and "
          "forage on the ground need it and pay for it. Controlling for mass, the fraction of "
          "ground foraging correlates +0.29 with leg residual and arboreal foraging −0.21. "
          "Water is the weak one, −0.06, because waterbirds are heterogeneous — wading herons and "
          "diving cormorants are both 'water'. That heterogeneity is the honest caveat here.")

# ============ 8 pipeline ============
sl = slide("Handing the law to the machine",
           "The prior does not choose the design. It chooses the region the generator is allowed to search.")
pic_fit(sl, f"{F}/en_F9_pipeline.png", .45, 1.45, W-.90, 4.45)
takeaway(sl, "A 1σ corridor, not a number — so the generator can still be told 'no' by the physics.")
notes(sl, "Important framing point for a mechanics audience: we are not copying a bird. The law "
          "gives a corridor of plausible leg lengths for a target mass; the generator proposes "
          "9-dimensional designs conditioned on a 5-dimensional operating point; and every "
          "candidate is checked by a full multibody drop simulation, not by a surrogate score. "
          "The biology narrows the search; the physics decides.")

# ============ 9 A0 ============
sl = slide("One page of engineering",
           "Three-segment leg, ankle and hip shock units, free knee — chosen by manufacturability, not by performance.")
pic_fit(sl, f"{F}/en_F8_a0.png", .35, 1.35, W-.70, 4.75)
takeaway(sl, "An architecture ablation showed all three candidates are equivalent within the "
             "seed-to-seed noise floor — so A0 wins on part count and on being able to buy the spring.", RED)
notes(sl, "Segments named after the bird: tarsometatarsus, tibiotarsus, femur. Two shock units, "
          "not three: we ablated the knee spring and the difference in peak deceleration was 3.60 "
          "vs 3.52 g, below the 2.9% we get from changing the random seed. Since performance does "
          "not distinguish the architectures, manufacturability does: A0 keeps a 36 mm lever arm on "
          "the femur, which turns a 1385 kN/m hip spring into a 616 kN/m one — the difference "
          "between a special order and a catalogue part.")

# ============ 10 b_eff ============
sl = slide("Where the machine disagrees with the bird",
           "Same joints, same 30 kg airframe, same touchdown speed. Only the scaling exponent differs.")
from PIL import Image as _Im
_iw, _ih = _Im.open(f"{A}/b_compare_en_last.png").size
_w = 8.55; _h = _w*_ih/_iw
if _h > 4.85: _h = 4.85; _w = _h*_iw/_ih
sl.shapes.add_movie(f"{A}/b_compare_en.mp4", I(.35+(8.55-_w)/2), I(1.30+(4.85-_h)/2),
                    I(_w), I(_h), poster_frame_image=f"{A}/b_compare_en_last.png",
                    mime_type="video/mp4")
tb(sl, 9.25, 1.45, 3.6, .45, "What the generator chose", 17, True, RED)
tb(sl, 9.25, 2.00, 3.7, 3.4,
   "Trained on drop simulations, the\n"
   "generator's own leg-length scaling\n"
   "settles at  b_eff = 0.238,  much\n"
   "flatter than biology's 0.365.\n\n"
   "At 30 kg that is 147 mm instead of\n"
   "184 mm: 29% less leg mass, and\n"
   "5 mm of stroke budget bought back,\n"
   "for 6% more peak deceleration.\n\n"
   "A bird cannot make that trade —\n"
   "it cannot fit a stiff, well-damped\n"
   "actuator into a bone. We can.",
   14, False, INK, space=1.30)
takeaway(sl, "Next:  add a rebound criterion to the feasibility test  ·  covariate-OU on the tree  ·  "
             "decide whether horizontal touchdown velocity enters the conditioning.")
notes(sl, "This is the punchline. We give the generator the biological corridor and let the "
          "simulation argue back. It ends up flatter than the bird — the machine can afford joint "
          "stiffness and damping that no tendon-and-bone leg can, so it does not need the extra "
          "length. The trade is 6% more peak g for 29% less leg mass, and both designs stay inside "
          "10 g and 24 mm. That number, 0.238 versus 0.365, is the clearest thing we have that this "
          "pipeline is doing something other than imitating nature. "
          "Open items: the feasibility test still has no rebound criterion, which blocks the "
          "hardware paper; and horizontal velocity is not yet a conditioning variable.")

out = f"{U}/07_09_2026_weekly_report_EN.pptx"
prs.save(out); print("→", out, os.path.getsize(out)//1024, "KB", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
